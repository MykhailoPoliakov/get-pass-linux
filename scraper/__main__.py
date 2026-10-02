import requests
import time
import json
import os
 
import openpyxl
from openpyxl import Workbook
 
with open("config.json") as f:
    config = json.load(f)
 
URL = config["url"]
EXCEL_FILE = "credentials.xlsx"
CHECK_INTERVAL = 5  # seconds between checks
 
 
def load_or_create_sheet():
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["User", "Password", "Date Changed"]) #type: ignore
    return wb, ws
 
 
def build_user_index(ws):
    # maps username -> row index, skipping header row
    index = {}
    for row in range(2, ws.max_row + 1):
        user = ws.cell(row=row, column=1).value
        if user:
            index[user] = row
    return index
 
 
def update_credential(ws, user_index, user, password, timestamp):
    if user in user_index:
        row = user_index[user]
        ws.cell(row=row, column=2, value=password)
        ws.cell(row=row, column=3, value=timestamp)
    else:
        ws.append([user, password, timestamp])
        user_index[user] = ws.max_row
 
 
def parse_content(content):
    """
    Expects one 'user:password' pair per line. Blank lines and lines
    without a colon are skipped. Adjust this function if your log's
    format is different (e.g. JSON, CSV, etc.).
    """
    pairs = []
    for line in content.splitlines():
        line = line.strip()
        if not line or ":" not in line:
            continue
        user, password = line.split(":", 1)
        pairs.append((user.strip(), password.strip()))
    return pairs
 
 
def watch_log():
    last_content = None
    wb, ws = load_or_create_sheet()
    user_index = build_user_index(ws)
 
    while True:
        try:
            response = requests.get(URL, timeout=10)
            current_content = response.text
 
            if current_content != last_content:
                timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
                pairs = parse_content(current_content)
 
                for user, password in pairs:
                    update_credential(ws, user_index, user, password, timestamp)
 
                wb.save(EXCEL_FILE)
 
                last_content = current_content
                print(f"[{time.strftime('%H:%M:%S')}] Change saved ({len(pairs)} entries).")
 
        except requests.RequestException as e:
            print(f"Error fetching log: {e}")
 
        time.sleep(CHECK_INTERVAL)
 
 
if __name__ == "__main__":
    watch_log()